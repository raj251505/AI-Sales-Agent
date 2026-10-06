import asyncio
import openpyxl
from datetime import datetime
from scraper import run_scraper
from scorer import score_all_leads, filter_actionable_leads
from telegram_bot import notify_batch, send_message, ApprovalHandler
from outreach import send_outreach
from crm import log_batch, update_lead_status, init_sheet
from scheduler import start_scheduler
from config import AGENT_NAME, AGENT_SERVICE

pending_leads: dict[str, dict] = {}

def load_leads_from_excel(filepath: str) -> list[dict]:
    wb = openpyxl.load_workbook(filepath)
    sheet_names = wb.sheetnames
    ws = None
    for name in sheet_names:
        if "No Website" in name or "Priority" in name:
            ws = wb[name]
            break
    if not ws:
        ws = wb.active

    headers = [cell.value for cell in ws[1]]
    leads = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        lead = dict(zip(headers, row))
        leads.append({
            "name": str(lead.get("Gym Name") or lead.get("Name") or ""),
            "phone": str(lead.get("Phone") or ""),
            "address": str(lead.get("Address") or ""),
            "rating": str(lead.get("Rating") or ""),
            "reviews": str(lead.get("Reviews") or ""),
            "website": str(lead.get("Website") or "—"),
            "email": str(lead.get("Email") or ""),
            "hours": str(lead.get("Hours") or ""),
            "maps_link": str(lead.get("Google Maps Link") or lead.get("Maps Link") or ""),
        })
        if leads[-1]["website"] == "—":
            leads[-1]["website"] = ""

    print(f"Loaded {len(leads)} leads from {filepath}")
    return leads

async def on_approve(lead_id: str):
    lead = pending_leads.get(lead_id)
    if not lead:
        await send_message(f"Lead {lead_id} not found")
        return

    print(f"Approved: {lead['name']}")
    await asyncio.to_thread(update_lead_status, lead.get("crm_id", lead_id), "Approved")

    await send_message(f"Writing personalized outreach for *{lead['name']}*...")
    await send_outreach(lead)

    await asyncio.to_thread(
        update_lead_status,
        lead.get("crm_id", lead_id),
        "Outreach Sent",
        {"outreach_sent": datetime.now().strftime("%Y-%m-%d %H:%M")}
    )

    await send_message(f"Outreach sent to {lead['name']}. I will follow up in 48hrs if no reply.")

async def on_reject(lead_id: str):
    lead = pending_leads.get(lead_id)
    if not lead:
        return

    print(f"Rejected: {lead['name']}")
    await asyncio.to_thread(update_lead_status, lead.get("crm_id", lead_id), "Rejected")

async def run_agent(excel_path: str = None, scrape_new: bool = False,
                    business_type: str = "gym", city: str = "Mumbai"):

    print("=" * 60)
    print(f"  {AGENT_NAME}'s AI Sales Agent")
    print(f"  Service: {AGENT_SERVICE}")
    print("=" * 60)

    await asyncio.to_thread(init_sheet)

    if scrape_new:
        print("\nStarting fresh scrape...")
        excel_path = await run_scraper(business_type, city)

    if excel_path:
        leads = load_leads_from_excel(excel_path)
    else:
        print("No leads source provided")
        return

    print("\nScoring leads...")
    scored_leads = score_all_leads(leads)
    actionable = filter_actionable_leads(scored_leads)

    print("\nLogging to Google Sheets CRM...")
    await asyncio.to_thread(log_batch, actionable)

    for i, lead in enumerate(actionable):
        lead_id = str(i)
        lead["_id"] = lead_id
        pending_leads[lead_id] = lead

    print("\nSending Telegram notifications...")
    await notify_batch(actionable)

    start_scheduler()

    print("\nListening for your approval on Telegram...")
    print("Reply /approve_ID or /reject_ID to any lead")
    print("Reply /status to see pipeline summary\n")

    handler = ApprovalHandler(on_approve=on_approve, on_reject=on_reject)
    await handler.start()

if __name__ == "__main__":
    print("\nAI Sales Agent Starting...\n")
    mode = input("Mode:\n1. Use existing Excel file\n2. Scrape fresh data\nChoice (1/2): ").strip()

    if mode == "1":
        excel_path = input("Path to Excel file: ").strip()
        asyncio.run(run_agent(excel_path=excel_path))

    elif mode == "2":
        business_type = input("Business type (e.g. gym, restaurant): ").strip()
        city = input("City (e.g. Mumbai, Pune): ").strip()
        asyncio.run(run_agent(scrape_new=True, business_type=business_type, city=city))

    else:
        print("Invalid choice")